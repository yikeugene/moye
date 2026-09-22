"""Create synthetic, offline DOCX/PPTX import fixtures; never opens Office.

Requires python-docx, python-pptx and Pillow (available in the Codex workspace
runtime). Run with an explicit output directory under artifacts, for example:
    python tools/generate-office-fixtures.py artifacts/office-import-qa/source
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGB
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def marker_image(path: Path) -> None:
    """Large, exact-color regions make image-loss checks unambiguous."""
    image = Image.new("RGB", (400, 240), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 199, 119), fill=(220, 40, 50))
    draw.rectangle((200, 0, 399, 119), fill=(35, 105, 225))
    draw.rectangle((0, 120, 199, 239), fill=(35, 165, 90))
    draw.rectangle((200, 120, 399, 239), fill=(245, 195, 45))
    image.save(path)


def word_fixture(path: Path, image_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = DocxInches(8.2677165354), DocxInches(11.692913386)
    section.top_margin = section.bottom_margin = DocxInches(.75)
    section.left_margin = section.right_margin = DocxInches(.75)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", DocxPt(12)
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "Microsoft JhengHei")

    document.add_heading("MOYE WORD FIXTURE — PAGE ONE", 0)
    document.add_paragraph("Synthetic mathematics lecture. No private notes or external resources.")
    document.add_paragraph("繁體中文課堂筆記：微積分與線性代數。")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Bold blue equation: ")
    run.bold = True
    run.font.color.rgb = DocxRGB(35, 105, 225)
    paragraph.add_run("f(x) = x² + 3x + 2")
    document.add_picture(str(image_path), width=DocxInches(3.5))
    document.add_paragraph("Image check: red / blue above green / yellow.")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Topic", "Value", "Review")):
        cell.text = text
    for values in (("Derivative", "2x + 3", "Week 1"), ("Integral", "x³/3", "Week 2")):
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text

    document.add_page_break()
    document.add_heading("MOYE WORD FIXTURE — PAGE TWO", 0)
    document.add_paragraph("Second page must remain a separate notebook page.")
    document.add_paragraph("第二頁：在匯入文件上加入手寫批註。")
    for text in ("First lecture point", "Second lecture point", "Third lecture point"):
        document.add_paragraph(text, style="List Bullet")
    document.add_picture(str(image_path), width=DocxInches(2.4))
    document.add_paragraph("END OF WORD FIXTURE")
    document.core_properties.title = "Synthetic Moye Word import fixture"
    document.core_properties.author = "Moye regression tests"
    document.save(path)


def slide_text(slide, text: str, x: float, y: float, width: float, height: float, size: int, color=(37, 51, 74)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Microsoft JhengHei"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor(*color)
    return box


def slide_fixture(path: Path, image_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.3333333333), Inches(7.5)
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_text(first, "MOYE POWERPOINT FIXTURE — SLIDE ONE", .6, .5, 12, .8, 28)
    slide_text(first, "Synthetic lecture slides\n繁體中文：矩陣與向量\nf(x) = x² + 3x + 2", .7, 1.6, 7.3, 2.8, 26)
    first.shapes.add_picture(str(image_path), Inches(8.6), Inches(1.6), width=Inches(3.6))
    shape = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.7), Inches(5.5), Inches(11.8), Inches(.7))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(35, 105, 225)
    shape.line.fill.background()
    slide_text(first, "BLUE FOOTER — IMAGE AND SHAPE CHECK", .9, 5.55, 11, .5, 20, (255, 255, 255))

    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second._element.set("show", "0")  # Hidden in slideshow, but still required in an imported notebook.
    slide_text(second, "MOYE POWERPOINT FIXTURE — SLIDE TWO", .6, .5, 12, .8, 28)
    slide_text(second, "Second slide must be a separate page.\n第二張投影片：保留文字、圖片及位置。", .7, 1.6, 11.5, 1.6, 26)
    shape = second.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.8), Inches(3.5), Inches(5.5), Inches(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(35, 165, 90)
    shape.line.fill.background()
    slide_text(second, "GREEN DIAGRAM", 1, 4.1, 5.1, .7, 24, (255, 255, 255))
    second.shapes.add_picture(str(image_path), Inches(8), Inches(3.4), width=Inches(3.5))
    slide_text(second, "END OF POWERPOINT FIXTURE", .7, 6.3, 11, .6, 18)
    jump = slide_text(first, "Jump to slide two (internal hyperlink)", .7, 6.5, 11, .6, 18, (35, 105, 225))
    jump.click_action.target_slide = second
    presentation.core_properties.title = "Synthetic Moye PowerPoint import fixture"
    presentation.core_properties.author = "Moye regression tests"
    presentation.save(path)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    output = parser.parse_args().output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    image_path = output / "four-color-marker.png"
    marker_image(image_path)
    word = output / "synthetic-lecture.docx"
    slides = output / "synthetic-slides.pptx"
    word_fixture(word, image_path)
    slide_fixture(slides, image_path)
    manifest = {
        "description": "Synthetic offline import fixtures; no user data or linked assets.",
        "files": [
            {"name": word.name, "pages": 2, "pageWidthDip": 793.700787, "pageHeightDip": 1122.519685,
             "textMarkers": ["MOYE WORD FIXTURE", "PAGE TWO", "END OF WORD FIXTURE"],
             "sha256": hashlib.sha256(word.read_bytes()).hexdigest()},
            {"name": slides.name, "pages": 2, "pageWidthDip": 1280, "pageHeightDip": 720, "hiddenSlideNumbers": [2],
             "textMarkers": ["MOYE POWERPOINT FIXTURE", "SLIDE TWO", "END OF POWERPOINT FIXTURE"],
             "sha256": hashlib.sha256(slides.read_bytes()).hexdigest()},
        ],
        "embeddedImageRgb": [[220, 40, 50], [35, 105, 225], [35, 165, 90], [245, 195, 45]],
    }
    (output / "expected.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Created synthetic DOCX/PPTX fixtures and expected.json in {output}")


if __name__ == "__main__":
    main()
